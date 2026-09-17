from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo

import psycopg
from openpyxl import load_workbook

from .download_dataset import sha256
from .settings import (
    DATASET_CITATION,
    DATASET_DOI,
    DATASET_LICENSE,
    DATASET_PAGE,
    DATASET_SLUG,
    DATASET_VERSION,
    WORKBOOK_PATH,
    database_url,
)

SPECIAL_CODES = re.compile(r"^(POST|D|M|DOT|BANK CHARGES|TEST\d*|AMAZONFEE|CRUK)$", re.IGNORECASE)
LONDON = ZoneInfo("Europe/London")


def clean_text(value: object, limit: int) -> str:
    return " ".join(str(value or "").replace("\x00", "").split())[:limit]


def parse_decimal(value: object) -> Decimal:
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError) as error:
        raise ValueError("decimal inválido") from error


def create_tracking(connection: psycopg.Connection, file_hash: str) -> tuple[str, str] | None:
    with connection.transaction():
        existing = connection.execute(
            "SELECT id::text,status FROM dataset_versions WHERE slug=%s AND file_sha256=%s",
            (DATASET_SLUG, file_hash),
        ).fetchone()
        if existing and existing[1] == "ready":
            print(f"A versão {existing[0]} já foi carregada; nenhuma duplicação foi criada.")
            return None
        dataset = connection.execute(
            """
            INSERT INTO dataset_versions(slug,version,source_url,doi,license,file_sha256,status)
            VALUES(%s,%s,%s,%s,%s,%s,'loading')
            ON CONFLICT(slug,version) DO UPDATE SET file_sha256=excluded.file_sha256,status='loading'
            RETURNING id::text
            """,
            (DATASET_SLUG, DATASET_VERSION, DATASET_PAGE, DATASET_DOI, DATASET_LICENSE, file_hash),
        ).fetchone()[0]
        sync = connection.execute(
            """INSERT INTO sync_runs(source,status,dataset_version_id,started_at)
               VALUES('uci-online-retail-ii','running',%s,now()) RETURNING id::text""",
            (dataset,),
        ).fetchone()[0]
        return dataset, sync


def stage_workbook(connection: psycopg.Connection) -> tuple[Counter, datetime, datetime]:
    counters: Counter = Counter()
    first_date: datetime | None = None
    last_date: datetime | None = None
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        with connection.cursor().copy(
            """COPY uci_import_stage(source_row_id,source_sheet,invoice_no,stock_code,description,
               invoice_at,invoice_date,quantity,unit_price,revenue,country,is_cancellation,source_hash)
               FROM STDIN"""
        ) as copy:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                headers = tuple(next(rows))
                expected = ("Invoice", "StockCode", "Description", "Quantity", "InvoiceDate", "Price", "Customer ID", "Country")
                if headers != expected:
                    raise RuntimeError(f"Cabeçalho inesperado em {sheet.title}: {headers}")
                for row_number, row in enumerate(rows, start=2):
                    counters["read"] += 1
                    invoice, stock_code, description, quantity, invoice_at, price, customer_id, country = row
                    if customer_id in (None, ""):
                        counters["missing_customer"] += 1
                    try:
                        invoice_text = clean_text(invoice, 40)
                        stock_text = clean_text(stock_code, 60).upper()
                        description_text = clean_text(description, 500)
                        country_text = clean_text(country, 100) or "Unknown"
                        if not invoice_text or not stock_text or not isinstance(invoice_at, datetime):
                            raise ValueError("identificador ou data ausente")
                        quantity_value = parse_decimal(quantity)
                        price_value = parse_decimal(price)
                        timestamp = invoice_at.replace(tzinfo=LONDON)
                        cancellation = invoice_text.upper().startswith("C") or quantity_value < 0
                        if cancellation:
                            counters["cancellation"] += 1
                        if not description_text:
                            counters["missing_description"] += 1
                        if price_value <= 0:
                            counters["nonpositive_price"] += 1
                        if quantity_value <= 0:
                            counters["nonpositive_quantity"] += 1
                        if SPECIAL_CODES.match(stock_text):
                            counters["non_product_code"] += 1
                        source_row_id = f"{sheet.title}:{row_number}"
                        fingerprint = "|".join((invoice_text, stock_text, description_text, str(quantity_value), timestamp.isoformat(), str(price_value), country_text))
                        copy.write_row((source_row_id, sheet.title, invoice_text, stock_text, description_text or None,
                                        timestamp, timestamp.date(), quantity_value, price_value,
                                        quantity_value * price_value, country_text, cancellation,
                                        hashlib.sha256(fingerprint.encode("utf-8")).hexdigest()))
                        first_date = timestamp if first_date is None or timestamp < first_date else first_date
                        last_date = timestamp if last_date is None or timestamp > last_date else last_date
                    except (ValueError, TypeError):
                        counters["malformed"] += 1
                    if counters["read"] % 100_000 == 0:
                        print(f"{counters['read']:,} linhas lidas...")
    finally:
        workbook.close()
    if first_date is None or last_date is None:
        raise RuntimeError("Nenhuma linha válida foi encontrada na planilha.")
    return counters, first_date, last_date


def persist(connection: psycopg.Connection, dataset_id: str, sync_id: str, counters: Counter,
            first_date: datetime, last_date: datetime) -> None:
    product_filter = """NOT is_cancellation AND quantity>0 AND unit_price>0 AND description IS NOT NULL
      AND stock_code !~* '^(POST|D|M|DOT|BANK CHARGES|TEST[0-9]*|AMAZONFEE|CRUK)$'"""
    connection.execute(
        f"""
        INSERT INTO products(external_id,sku,name,unit,currency,sale_price,active,source)
        SELECT %s||':'||stock_code, stock_code,
               (array_agg(description ORDER BY invoice_at DESC))[1], 'UN', 'GBP',
               (array_agg(unit_price ORDER BY invoice_at DESC))[1], true, 'uci_online_retail'
        FROM uci_import_stage WHERE {product_filter}
        GROUP BY stock_code
        ON CONFLICT(external_id) DO UPDATE SET name=excluded.name,sale_price=excluded.sale_price,
          currency='GBP',active=true,source='uci_online_retail',updated_at=now(),version=products.version+1
        """,
        (DATASET_SLUG,),
    )
    connection.execute(
        """
        INSERT INTO retail_sales_facts(dataset_version_id,invoice_date,source_row_id,source_sheet,invoice_no,
          stock_code,description,product_id,invoice_at,quantity,unit_price,revenue,country,is_cancellation,is_duplicate,source_hash)
        SELECT %s,s.invoice_date,s.source_row_id,s.source_sheet,s.invoice_no,s.stock_code,s.description,p.id,
          s.invoice_at,s.quantity,s.unit_price,s.revenue,s.country,s.is_cancellation,
          row_number() OVER(PARTITION BY s.source_hash ORDER BY s.source_sheet,s.source_row_id)>1,s.source_hash
        FROM uci_import_stage s LEFT JOIN products p ON p.external_id=%s||':'||s.stock_code
        ON CONFLICT DO NOTHING
        """,
        (dataset_id, DATASET_SLUG),
    )
    connection.execute(
        """
        INSERT INTO sales_orders(external_id,order_number,sold_at,status,channel,total_amount,source_updated_at)
        SELECT %s||':'||source_sheet||':'||invoice_no,invoice_no,min(invoice_at),'completed',country,
          round(sum(revenue),2),max(invoice_at)
        FROM retail_sales_facts
        WHERE dataset_version_id=%s AND NOT is_cancellation AND NOT is_duplicate AND product_id IS NOT NULL
          AND quantity>0 AND unit_price>0
        GROUP BY source_sheet,invoice_no,country
        ON CONFLICT(external_id) DO UPDATE SET total_amount=excluded.total_amount,source_updated_at=excluded.source_updated_at,updated_at=now()
        """,
        (DATASET_SLUG, dataset_id),
    )
    connection.execute(
        """
        WITH valid AS (
          SELECT f.*,row_number() OVER(PARTITION BY f.source_sheet,f.invoice_no ORDER BY f.source_row_id) AS line_number
          FROM retail_sales_facts f WHERE f.dataset_version_id=%s AND NOT f.is_cancellation AND NOT f.is_duplicate
            AND f.product_id IS NOT NULL AND f.quantity>0 AND f.unit_price>0
        )
        INSERT INTO sales_order_items(order_id,product_id,line_number,quantity,unit_price,discount_amount,total_amount)
        SELECT o.id,v.product_id,v.line_number,v.quantity,v.unit_price,0,round(v.revenue,2)
        FROM valid v JOIN sales_orders o ON o.external_id=%s||':'||v.source_sheet||':'||v.invoice_no
        ON CONFLICT(order_id,line_number) DO UPDATE SET product_id=excluded.product_id,quantity=excluded.quantity,
          unit_price=excluded.unit_price,total_amount=excluded.total_amount
        """,
        (dataset_id, DATASET_SLUG),
    )
    connection.execute(
        """
        INSERT INTO daily_product_demand(product_id,demand_date,units_sold,gross_revenue,orders_count,was_stockout,dataset_version_id)
        SELECT product_id,invoice_date,sum(quantity),round(sum(revenue),2),count(DISTINCT source_sheet||':'||invoice_no),NULL,%s
        FROM retail_sales_facts WHERE dataset_version_id=%s AND NOT is_cancellation AND NOT is_duplicate
          AND product_id IS NOT NULL AND quantity>0 AND unit_price>0
        GROUP BY product_id,invoice_date
        ON CONFLICT(product_id,demand_date) DO UPDATE SET units_sold=excluded.units_sold,
          gross_revenue=excluded.gross_revenue,orders_count=excluded.orders_count,
          dataset_version_id=excluded.dataset_version_id,calculated_at=now()
        """,
        (dataset_id, dataset_id),
    )
    # A fonte possui vendas, mas não saldos. Criamos um depósito de simulação
    # separado e um saldo reprodutível para permitir demonstrar o módulo de
    # estoque sem fingir que se trata de uma medição física da varejista.
    warehouse_id = connection.execute(
        """INSERT INTO warehouses(external_id,name,active,source)
           VALUES('uci-demo-warehouse','Depósito simulado (UCI)',true,'dataset')
           ON CONFLICT(external_id) DO UPDATE SET name=excluded.name,active=true,source='dataset',updated_at=now()
           RETURNING id"""
    ).fetchone()[0]
    connection.execute(
        """WITH anchor AS (SELECT max(demand_date) AS day FROM daily_product_demand WHERE dataset_version_id=%s),
        recent AS (
          SELECT d.product_id, sum(d.units_sold) / 30.0 AS daily_average
          FROM daily_product_demand d CROSS JOIN anchor a
          WHERE d.dataset_version_id=%s AND d.demand_date BETWEEN a.day - 29 AND a.day
          GROUP BY d.product_id
        )
        INSERT INTO inventory_levels(product_id,warehouse_id,on_hand_quantity,reserved_quantity,measured_at)
        SELECT p.id,%s,round(GREATEST(COALESCE(r.daily_average,0) * 30,0),4),0,%s
        FROM products p LEFT JOIN recent r ON r.product_id=p.id
        WHERE p.source='uci_online_retail'
        ON CONFLICT(product_id,warehouse_id) DO UPDATE SET
          on_hand_quantity=excluded.on_hand_quantity,reserved_quantity=0,
          measured_at=excluded.measured_at,updated_at=now()""",
        (dataset_id, dataset_id, warehouse_id, last_date),
    )
    connection.execute(
        """WITH anchor AS (SELECT max(demand_date) AS day FROM daily_product_demand WHERE dataset_version_id=%s),
        recent AS (
          SELECT d.product_id, sum(d.units_sold) / 30.0 AS daily_average
          FROM daily_product_demand d CROSS JOIN anchor a
          WHERE d.dataset_version_id=%s AND d.demand_date BETWEEN a.day - 29 AND a.day
          GROUP BY d.product_id
        )
        UPDATE products p SET
          minimum_stock=round(GREATEST(COALESCE(r.daily_average,0) * 7,0),4),
          lead_time_days=7,
          cost_price=round(COALESCE(p.sale_price,0) * 0.60,4),
          updated_at=now()
        FROM recent r WHERE p.id=r.product_id AND p.source='uci_online_retail'""",
        (dataset_id, dataset_id),
    )
    counts = connection.execute(
        """SELECT count(*)::bigint,
          count(*) FILTER (WHERE is_duplicate),
          count(*) FILTER (WHERE NOT is_cancellation AND NOT is_duplicate AND product_id IS NOT NULL AND quantity>0 AND unit_price>0)
          FROM retail_sales_facts WHERE dataset_version_id=%s""",
        (dataset_id,),
    ).fetchone()
    counters["staged"] = counts[0]
    counters["duplicate"] = counts[1]
    counters["accepted"] = counts[2]
    counters["rejected"] = counters["read"] - counters["accepted"]
    rules = [
        ("exact_duplicates", counters["accepted"], counters["duplicate"]),
        ("cancellations", counters["accepted"], counters["cancellation"]),
        ("missing_description", counters["accepted"], counters["missing_description"]),
        ("nonpositive_price", counters["accepted"], counters["nonpositive_price"]),
        ("non_product_code", counters["accepted"], counters["non_product_code"]),
        ("malformed_rows", counters["accepted"], counters["malformed"]),
    ]
    with connection.cursor() as cursor:
        cursor.executemany(
            """INSERT INTO data_quality_results(sync_run_id,rule_code,severity,passed_count,failed_count,sample)
               VALUES(%s,%s,%s,%s,%s,'[]'::jsonb)
               ON CONFLICT(sync_run_id,rule_code) DO UPDATE SET passed_count=excluded.passed_count,
                 failed_count=excluded.failed_count,checked_at=now()""",
            [(sync_id, code, "warning" if failed else "info", passed, failed) for code, passed, failed in rules],
        )
    summary = dict(counters)
    summary["citation"] = DATASET_CITATION
    connection.execute(
        """UPDATE dataset_versions SET status='ready',records_read=%s,records_accepted=%s,records_rejected=%s,
          period_started_on=%s,period_ended_on=%s,quality_summary=%s::jsonb,imported_at=now() WHERE id=%s""",
        (counters["read"], counters["accepted"], counters["rejected"], first_date.date(), last_date.date(), json.dumps(summary), dataset_id),
    )
    connection.execute(
        """UPDATE sync_runs SET status='succeeded',cursor_started_at=%s,cursor_ended_at=%s,records_read=%s,
          records_inserted=%s,records_rejected=%s,finished_at=now() WHERE id=%s""",
        (first_date, last_date, counters["read"], counters["accepted"], counters["rejected"], sync_id),
    )


def mark_failed(connection: psycopg.Connection, dataset_id: str, sync_id: str, error: Exception) -> None:
    message = f"{type(error).__name__}: {error}"[:1000]
    with connection.transaction():
        connection.execute("UPDATE dataset_versions SET status='failed' WHERE id=%s", (dataset_id,))
        connection.execute("UPDATE sync_runs SET status='failed',error_message=%s,finished_at=now() WHERE id=%s", (message, sync_id))


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise RuntimeError("Execute npm run data:download antes da carga.")
    file_hash = sha256(WORKBOOK_PATH)
    with psycopg.connect(database_url(), autocommit=False) as connection:
        tracking = create_tracking(connection, file_hash)
        if tracking is None:
            return
        dataset_id, sync_id = tracking
        try:
            with connection.transaction():
                connection.execute("SET LOCAL statement_timeout = 0")
                connection.execute("SET LOCAL lock_timeout = '10s'")
                connection.execute("""CREATE TEMP TABLE uci_import_stage(
                  source_row_id text,source_sheet text,invoice_no text,stock_code text,description text,
                  invoice_at timestamptz,invoice_date date,quantity numeric(14,4),unit_price numeric(14,4),
                  revenue numeric(16,4),country text,is_cancellation boolean,source_hash char(64)) ON COMMIT DROP""")
                counters, first_date, last_date = stage_workbook(connection)
                persist(connection, dataset_id, sync_id, counters, first_date, last_date)
            print(f"Carga concluída: {counters['accepted']:,} linhas aceitas de {counters['read']:,}.")
        except Exception as error:
            mark_failed(connection, dataset_id, sync_id, error)
            raise


if __name__ == "__main__":
    main()
